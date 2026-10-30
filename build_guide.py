from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
import pypandoc

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "src"
COMBINED = ROOT / "combined.md"
DOCX = ROOT / "dbt_from_zero_to_expert.docx"
HTML = ROOT / "dbt_from_zero_to_expert.html"
PDF = ROOT / "dbt_from_zero_to_expert.pdf"
FILTER = ROOT / "filters" / "heading_classes.lua"
CSS = ROOT / "guide.css"

PARTS_BEFORE_DATABRICKS = [
    "part03_jinja.md",
    "part04_testing.md",
    "part05_incremental.md",
]
PARTS_AFTER_DATABRICKS = [
    "part07_advanced.md",
    "part09_semantic_ai.md",
    "part10_capstone.md",
    "part11_interview.md",
]


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def assemble() -> str:
    front = read(SRC / "part00_front.md")
    marker = "\n# PART III - REUSABILITY"
    if marker not in front:
        raise SystemExit("Could not find Part III marker in part00_front.md")
    front = front.split(marker, 1)[0].rstrip() + "\n"

    databricks = read(SRC / "part06_databricks.md")
    databricks_marker = "\n# PART VII - ADVANCED ENGINEERING"
    if databricks_marker not in databricks:
        raise SystemExit("Could not find Part VII marker in part06_databricks.md")
    databricks = databricks.split(databricks_marker, 1)[0].rstrip() + "\n"

    chunks = (
        [front]
        + [read(SRC / name) for name in PARTS_BEFORE_DATABRICKS]
        + [databricks]
        + [read(SRC / name) for name in PARTS_AFTER_DATABRICKS]
    )
    combined = "\n\n".join(chunk.strip() + "\n" for chunk in chunks)
    COMBINED.write_text(combined, encoding="utf-8")
    return combined


def validate(combined: str) -> None:
    chapters = [int(n) for n in re.findall(r"^# (\d+)\. ", combined, flags=re.M)]
    if chapters != list(range(1, 21)):
        raise SystemExit(f"Chapter sequence is wrong: {chapters}")
    parts = re.findall(r"^# PART ([IVX]+) - ", combined, flags=re.M)
    if parts != ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]:
        raise SystemExit(f"Part sequence is wrong: {parts}")
    if combined.count("# PART ") != 10:
        raise SystemExit("Duplicate or missing part headings")
    if len(chapters) != len(set(chapters)):
        raise SystemExit("Duplicate chapter numbers")
    print(f"Validated: {len(parts)} parts, {len(chapters)} chapters, {len(combined):,} chars")


def pandoc_cmd(output: Path, extra: list[str]) -> list[str]:
    return [
        pypandoc.get_pandoc_path(),
        str(COMBINED),
        "--from=markdown",
        "--standalone",
        "--highlight-style=tango",
        "--lua-filter",
        str(FILTER),
        "--resource-path",
        str(SRC),
        "-o",
        str(output),
        *extra,
    ]


def build_docx() -> None:
    subprocess.run(pandoc_cmd(DOCX, []), check=True, cwd=str(ROOT))
    style_docx()
    print(f"DOCX written: {DOCX} ({DOCX.stat().st_size:,} bytes)")


def build_html() -> None:
    subprocess.run(
        pandoc_cmd(
            HTML,
            [
                "--css=guide.css",
                "--metadata=lang:en",
                "--wrap=none",
            ],
        ),
        check=True,
        cwd=str(ROOT),
    )
    print(f"HTML written: {HTML} ({HTML.stat().st_size:,} bytes)")


def build_pdf() -> None:
    edge_candidates = [
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    ]
    browser = next((p for p in edge_candidates if p.exists()), None)
    if browser is None:
        raise SystemExit("No Edge/Chrome executable found for PDF rendering")

    profile = Path(r"C:\Users\THINKB~1\AppData\Local\Temp\opencode\guide-edge-profile")
    profile.mkdir(parents=True, exist_ok=True)
    cmd = [
        str(browser),
        "--headless=new",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--user-data-dir={profile}",
        "--virtual-time-budget=20000",
        f"--print-to-pdf={PDF}",
        HTML.as_uri(),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=240)
    if not PDF.exists() or PDF.stat().st_size < 100_000:
        print(result.stdout)
        print(result.stderr)
        raise SystemExit(f"PDF was not produced correctly: {PDF}")
    print(f"PDF written: {PDF} ({PDF.stat().st_size:,} bytes)")


def stamp_pdf() -> None:
    doc = pymupdf.open(PDF)
    total = doc.page_count
    for index, page in enumerate(doc):
        if index == 0:
            continue
        rect = page.rect
        text = f"dbt from Zero to Expert  |  Page {index + 1} of {total}"
        page.insert_textbox(
            pymupdf.Rect(0, rect.height - 34, rect.width, rect.height - 14),
            text,
            fontsize=7.5,
            fontname="helv",
            color=(0.45, 0.45, 0.45),
            align=pymupdf.TEXT_ALIGN_CENTER,
        )

    headings: list[list] = []
    for line in read(COMBINED).splitlines():
        if line.startswith("# PART "):
            headings.append([1, line[2:].strip()])
        elif re.match(r"^# \d+\. ", line):
            headings.append([2, line[2:].strip()])

    pages_text = [re.sub(r"\s+", " ", page.get_text("text")) for page in doc]
    toc: list[list] = []
    cursor = 4
    for level, title in headings:
        for index in range(cursor, total):
            probe = re.sub(r"\s+", " ", title)[:40]
            if probe in pages_text[index]:
                toc.append([level, title, index + 1])
                cursor = index
                break
    if len(toc) == len(headings):
        doc.set_toc(toc)
        print(f"PDF bookmarks added: {len(toc)} entries")
    else:
        print(f"PDF bookmarks skipped: matched {len(toc)} of {len(headings)} headings")

    doc.set_metadata(
        {
            "title": "dbt from Zero to Expert",
            "author": "dbt from Zero to Expert",
            "subject": "A hands-on learning guide for VS Code and Databricks",
            "keywords": "dbt, Databricks, analytics engineering, SQL, data modeling",
            "creator": "Pandoc + Edge",
            "producer": "PyMuPDF",
        }
    )
    temp = PDF.with_suffix(".stamped.pdf")
    doc.save(temp, garbage=4, deflate=True)
    doc.close()
    temp.replace(PDF)
    print(f"PDF stamped with page numbers: {PDF} ({PDF.stat().st_size:,} bytes)")


def shade(element, fill: str) -> None:
    ppr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(
        shd,
        "w:tabs",
        "w:suppressAutoHyphens",
        "w:kinsoku",
        "w:wordWrap",
        "w:overflowPunct",
        "w:topLinePunct",
        "w:autoSpaceDE",
        "w:autoSpaceDN",
        "w:bidi",
        "w:adjustRightInd",
        "w:snapToGrid",
        "w:spacing",
        "w:ind",
        "w:contextualSpacing",
        "w:mirrorIndents",
        "w:suppressOverlap",
        "w:jc",
        "w:textDirection",
        "w:textAlignment",
        "w:textboxTightWrap",
        "w:outlineLvl",
        "w:divId",
        "w:cnfStyle",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )


def style_docx() -> None:
    doc = Document(str(DOCX))

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(22)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(8)
    heading1.paragraph_format.keep_with_next = True

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True

    heading3 = doc.styles["Heading 3"]
    heading3.font.name = "Calibri"
    heading3.font.size = Pt(11.5)
    heading3.font.bold = True
    heading3.font.color.rgb = RGBColor(0x2E, 0x5C, 0x8A)
    heading3.paragraph_format.space_before = Pt(9)
    heading3.paragraph_format.space_after = Pt(3)
    heading3.paragraph_format.keep_with_next = True

    for name, size, color in [
        ("Title", 30, RGBColor(0x1F, 0x4E, 0x79)),
        ("Subtitle", 16, RGBColor(0x2E, 0x5C, 0x8A)),
        ("Author", 11, RGBColor(0x33, 0x33, 0x33)),
        ("Date", 10, RGBColor(0x66, 0x66, 0x66)),
    ]:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        if name == "Title":
            style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    try:
        code_style = doc.styles["Source Code"]
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8.5)
        code_style.paragraph_format.space_before = Pt(0)
        code_style.paragraph_format.space_after = Pt(0)
        code_style.paragraph_format.line_spacing = 1.0
    except KeyError:
        pass

    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()

        if style_name == "Title":
            paragraph.paragraph_format.space_before = Pt(150)
            paragraph.paragraph_format.space_after = Pt(12)
        elif style_name == "Heading 1":
            if re.match(r"^PART [IVX]+ - ", text):
                paragraph.paragraph_format.page_break_before = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade(paragraph._p, "1F4E79")
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(24)
                    run.font.bold = True
            elif text == "How to Use This Guide":
                paragraph.paragraph_format.page_break_before = True
            else:
                paragraph.paragraph_format.page_break_before = False
                paragraph.paragraph_format.space_before = Pt(0)
        elif style_name == "Source Code":
            shade(paragraph._p, "F6F8FA")
            paragraph.paragraph_format.left_indent = Pt(4)
            paragraph.paragraph_format.right_indent = Pt(4)

    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "B8C4D0")
            borders.append(el)
        tbl_pr.insert_element_before(
            borders,
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
            "w:tblPrChange",
        )

        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            cant_split.set(qn("w:val"), "true")
            tr_pr.append(cant_split)
            if row_index == 0:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)

            for cell in row.cells:
                if row_index == 0:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "DBE5F1")
                    tc_pr.insert_element_before(
                        shd,
                        "w:noWrap",
                        "w:tcMar",
                        "w:textDirection",
                        "w:tcFitText",
                        "w:vAlign",
                        "w:hideMark",
                        "w:headers",
                        "w:cellIns",
                        "w:cellDel",
                        "w:cellMerge",
                        "w:tcPrChange",
                    )

                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.font.bold = True

    footer = doc.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = footer_paragraph.add_run("dbt from Zero to Expert  |  Page ")
    prefix.font.size = Pt(8)
    prefix.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    footer_paragraph._p.append(fld)

    props = doc.core_properties
    props.title = "dbt from Zero to Expert"
    props.subject = "A hands-on learning guide for VS Code and Databricks"
    props.author = "dbt from Zero to Expert"
    props.keywords = "dbt, Databricks, analytics engineering, SQL, data modeling"
    props.comments = "Edition 1.0, October 2026. dbt Core 1.12 and dbt v2 aware."

    doc.save(str(DOCX))


def verify() -> None:
    doc = Document(str(DOCX))
    h1 = [p.text.strip() for p in doc.paragraphs if p.style.name == "Heading 1"]
    parts = [t for t in h1 if t.startswith("PART ")]
    chapters = [t for t in h1 if re.match(r"^\d+\. ", t)]
    tables = len(doc.tables)
    print(f"DOCX check: paragraphs={len(doc.paragraphs):,} tables={tables} parts={len(parts)} chapters={len(chapters)}")
    if len(parts) != 10 or len(chapters) != 20:
        raise SystemExit("DOCX heading verification failed")

    reader = PdfReader(str(PDF))
    pages = len(reader.pages)
    first = " ".join((reader.pages[i].extract_text() or "") for i in range(min(4, pages)))
    full = " ".join((reader.pages[i].extract_text() or "") for i in range(pages))
    compact = re.sub(r"\s+", " ", full)
    print(f"PDF check: pages={pages} first_page_chars={len(first)} full_chars={len(compact):,}")
    for needle in ["PART X - MASTERY", "20. Cheat Sheets", "Interview Questions"]:
        if needle not in compact:
            raise SystemExit(f"PDF verification failed: missing {needle!r}")
    if not PDF.exists() or PDF.stat().st_size < 100_000:
        raise SystemExit("PDF file is missing or too small")


def main() -> None:
    combined = assemble()
    validate(combined)
    build_docx()
    build_html()
    build_pdf()
    stamp_pdf()
    verify()
    print("Build complete.")


if __name__ == "__main__":
    main()

